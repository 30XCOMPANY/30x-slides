"""
30x Slides — Upload PPTX, get interactive HTML slides with Voice Bot.
Stack: Flask + SocketIO + LibreOffice + pymupdf + Anthropic + Fish Audio

[INPUT]: PPTX file upload and viewer talk requests with optional slide context
[OUTPUT]: Interactive slide viewer with voice bot + downloadable HTML
[POS]: Application entry point, orchestrates upload, conversion, and service wiring
[PROTOCOL]: 变更时更新此头部，然后检查 AGENTS.md
"""

import os
import re
import json
import uuid
import base64
import subprocess
import urllib.request
import tempfile
import time
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from dotenv import load_dotenv

load_dotenv(dotenv_path=Path(__file__).with_name(".env"), override=False)
from flask import (
    Flask, request, redirect, url_for,
    render_template, send_from_directory, jsonify, Response,
)
from flask_socketio import SocketIO

import fitz  # pymupdf
from services.chat import build_talk_response
from services.config import log_boot_env, validate_required_env
from services.decks import extract_slide_content, load_deck_content, save_deck_content
from services.tts import stream_fish_audio, synthesize_fish_audio

app = Flask(__name__)
app.config["SECRET_KEY"] = "30x-slides-secret"
app.config["OUTPUT_FOLDER"] = os.path.join(os.path.dirname(__file__), "output")
app.config["MAX_CONTENT_LENGTH"] = 100 * 1024 * 1024  # 100MB

socketio = SocketIO(app, cors_allowed_origins="*", async_mode="threading")

# ---- 字体目录 (macOS vs Linux) ----
import platform
if platform.system() == "Darwin":
    FONT_DIR = Path.home() / "Library" / "Fonts"
else:
    FONT_DIR = Path.home() / ".local" / "share" / "fonts"
FONT_DIR.mkdir(parents=True, exist_ok=True)

log_boot_env()
validate_required_env()




# ============================================================
# 字体工具
# ============================================================
def get_pptx_fonts(pptx_path):
    """从 PPTX theme XML + slide XML 提取所有字体名称"""
    fonts = set()
    try:
        import zipfile as zf
        from lxml import etree
        with zf.ZipFile(pptx_path) as z:
            for name in z.namelist():
                if name.endswith(".xml"):
                    try:
                        tree = etree.parse(z.open(name))
                        for elem in tree.iter():
                            tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
                            if tag in ("latin", "ea", "cs", "rFont"):
                                tf = elem.get("typeface") or elem.get("val")
                                if tf and not tf.startswith("+") and tf not in ("", "Calibri"):
                                    fonts.add(tf)
                    except Exception:
                        pass
    except Exception:
        pass

    try:
        from pptx import Presentation
        prs = Presentation(pptx_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            if run.font.name:
                                fonts.add(run.font.name)
    except Exception:
        pass

    return fonts


def get_system_fonts():
    fonts = set()
    try:
        result = subprocess.run(
            ["fc-list", "--format", "%{family}\n"],
            capture_output=True, text=True,
        )
        for line in result.stdout.splitlines():
            name = line.strip().split(",")[0].strip()
            if name:
                fonts.add(name)
    except Exception:
        pass
    return sorted(fonts)


def is_font_installed(font_name):
    try:
        result = subprocess.run(["fc-list"], capture_output=True, text=True)
        normalized = font_name.lower().replace(" ", "")
        for line in result.stdout.splitlines():
            if normalized in line.lower().replace(" ", ""):
                return True
    except Exception:
        pass
    return False


def try_install_font(font_name):
    """只装 400 + 700 两个权重，够用且快"""
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", font_name).strip("-").lower()
    weights = {"400": "latin-400-normal", "700": "latin-700-normal"}
    installed = False
    for wk, ws in weights.items():
        url = f"https://cdn.jsdelivr.net/fontsource/fonts/{slug}@latest/{ws}.ttf"
        dest = FONT_DIR / f"{slug}-{wk}.ttf"
        if dest.exists():
            installed = True
            continue
        try:
            urllib.request.urlretrieve(url, str(dest))
            with open(dest, "rb") as f:
                header = f.read(4)
            if header[:4] in (b"\x00\x01\x00\x00", b"true", b"OTTO"):
                installed = True
            else:
                dest.unlink()
        except Exception:
            if dest.exists():
                dest.unlink()
    return installed


def ensure_fonts(pptx_path):
    fonts = get_pptx_fonts(pptx_path)
    report = {"found": [], "installed": [], "missing": []}
    for font in sorted(fonts):
        if is_font_installed(font):
            report["found"].append(font)
        elif try_install_font(font):
            report["installed"].append(font)
        else:
            report["missing"].append(font)
    if report["installed"]:
        subprocess.run(["fc-cache", "-f"], capture_output=True)
    print(f"[FONTS] found={len(report['found'])} installed={len(report['installed'])} missing={len(report['missing'])}", flush=True)
    return report


# ============================================================
# 内容提取
# ============================================================
def extract_slide_content(pptx_path):
    """提取每页 slide 的文字内容 + speaker notes"""
    slides_content = []
    try:
        from pptx import Presentation
        prs = Presentation(pptx_path)
        for slide in prs.slides:
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = para.text.strip()
                        if line:
                            texts.append(line)

            notes = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()

            slides_content.append({
                "slide": len(slides_content) + 1,
                "text": texts,
                "notes": notes,
            })
    except Exception as e:
        print(f"[extract] error: {e}")

    return slides_content


# ============================================================
# 转换管线
# ============================================================
def convert_pptx_to_pdf(pptx_path, deck_dir):
    pdf_path = os.path.join(deck_dir, "deck.pdf")
    soffice_profile = Path(tempfile.mkdtemp(prefix="soffice-profile-", dir=deck_dir))
    profile_uri = soffice_profile.resolve().as_uri()

    print(f"[CONVERT] LibreOffice starting: {pptx_path}", flush=True)
    result = subprocess.run(
        [
            "soffice",
            "--headless",
            f"-env:UserInstallation={profile_uri}",
            "--convert-to",
            "pdf",
            pptx_path,
            "--outdir",
            deck_dir,
        ],
        capture_output=True,
        text=True,
        timeout=120,
    )
    print(f"[CONVERT] LibreOffice done, rc={result.returncode}", flush=True)
    if result.returncode != 0:
        print(f"[CONVERT ERROR] {result.stderr}", flush=True)
        raise RuntimeError(f"LibreOffice failed: {result.stderr}")
    return pdf_path


def render_pdf_page(pdf_path, page_index, target_width=1920, target_height=1080):
    doc = fitz.open(pdf_path)
    try:
        page = doc[page_index]
        matrix = fitz.Matrix(target_width / page.rect.width, target_height / page.rect.height)
        pix = page.get_pixmap(matrix=matrix, alpha=False)
        output_path = os.path.join(os.path.dirname(pdf_path), f"slide-{page_index + 1}.png")
        pix.save(output_path)
        return output_path
    finally:
        doc.close()


def render_pdf_pages(pdf_path):
    with fitz.open(pdf_path) as doc:
        total = len(doc)

    workers = max(1, min(4, (os.cpu_count() or 1)))
    with ThreadPoolExecutor(max_workers=workers) as executor:
        list(executor.map(lambda idx: render_pdf_page(pdf_path, idx), range(total)))
    return total


def convert_deck(deck_dir):
    pptx_path = os.path.join(deck_dir, "deck.pptx")
    started_at = time.perf_counter()
    pdf_path = convert_pptx_to_pdf(pptx_path, deck_dir)

    with ThreadPoolExecutor(max_workers=2) as executor:
        raster_future = executor.submit(render_pdf_pages, pdf_path)
        content_future = executor.submit(extract_slide_content, pptx_path)
        total = raster_future.result()
        content = content_future.result()

    with open(os.path.join(deck_dir, "meta.txt"), "w") as f:
        f.write(str(total))
    save_deck_content(app.config["OUTPUT_FOLDER"], os.path.basename(deck_dir), content)

    elapsed = time.perf_counter() - started_at
    print(f"[CONVERT] deck={os.path.basename(deck_dir)} slides={total} elapsed={elapsed:.2f}s", flush=True)
    return total


def apply_font_replacements(pptx_path, replacements, weight_map=None):
    """替换 PPTX XML 中的字体名称，可选调整粗细"""
    import zipfile, shutil
    from lxml import etree

    tmp = pptx_path + ".tmp"
    with zipfile.ZipFile(pptx_path, "r") as zin, \
         zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.endswith(".xml") or item.filename.endswith(".rels"):
                text = data.decode("utf-8")
                for old, new in replacements.items():
                    text = text.replace(f'typeface="{old}"', f'typeface="{new}"')
                    text = text.replace(f'val="{old}"', f'val="{new}"')

                if weight_map and item.filename.startswith("ppt/slides/"):
                    try:
                        root = etree.fromstring(text.encode("utf-8"))
                        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                        for orig_font, weight in weight_map.items():
                            new_font = replacements.get(orig_font, orig_font)
                            w = int(weight)
                            for rpr in root.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr"):
                                latin = rpr.find("a:latin", ns)
                                if latin is not None and latin.get("typeface") == new_font:
                                    if w >= 600:
                                        rpr.set("b", "1")
                                    else:
                                        rpr.attrib.pop("b", None)
                        text = etree.tostring(root, xml_declaration=True,
                                              encoding="UTF-8", standalone=True).decode("utf-8")
                    except Exception:
                        pass

                data = text.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, pptx_path)


# ============================================================
# HTTP API — 对话式语音导览 (比 WebSocket 更稳定)
# ============================================================
@app.route("/api/talk/<deck_id>", methods=["POST"])
def api_talk(deck_id):
    """用户说话 → 对话服务 → 返回 JSON {text, nav}"""
    return jsonify(build_talk_response(deck_id, request.json or {}, app.config["OUTPUT_FOLDER"]))


@app.route("/api/tts", methods=["POST"])
def api_tts():
    """Fish Audio TTS — 返回 base64 mp3"""
    data = request.json or {}
    text = data.get("text", "").strip()
    try:
        return jsonify({"audio": synthesize_fish_audio(text)})
    except Exception as exc:
        print(f"[TTS ERROR] {exc}", flush=True)
        return jsonify({"audio": ""})


@app.route("/api/tts/stream")
def api_tts_stream():
    """Fish Audio streaming proxy — progressive MP3 bytes for immediate playback."""
    text = (request.args.get("text") or "").strip()
    if not text:
        return ("", 204)

    try:
        return Response(
            stream_fish_audio(text),
            mimetype="audio/mpeg",
            headers={
                "Cache-Control": "no-store",
                "X-Accel-Buffering": "no",
            },
            direct_passthrough=True,
        )
    except Exception as exc:
        print(f"[TTS STREAM ERROR] {exc}", flush=True)
        return ("", 204)


# ============================================================
# HTTP 路由
# ============================================================
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/upload", methods=["POST"])
def upload():
    file = request.files.get("file")
    if not file or not file.filename.endswith(".pptx"):
        return "Please upload a .pptx file", 400

    deck_id = uuid.uuid4().hex[:10]
    deck_dir = os.path.join(app.config["OUTPUT_FOLDER"], deck_id)
    os.makedirs(deck_dir, exist_ok=True)

    pptx_path = os.path.join(deck_dir, "deck.pptx")
    file.save(pptx_path)

    try:
        font_report = ensure_fonts(pptx_path)
        with open(os.path.join(deck_dir, "fonts.txt"), "w") as f:
            for cat in ("found", "installed", "missing"):
                for name in font_report[cat]:
                    f.write(f"{cat}: {name}\n")

        convert_deck(deck_dir)
        return redirect(url_for("viewer", deck_id=deck_id))
    except Exception as e:
        print(f"[UPLOAD ERROR] {e}", flush=True)
        return f"Conversion failed: {e}", 500


@app.route("/view/<deck_id>")
def viewer(deck_id):
    deck_dir = os.path.join(app.config["OUTPUT_FOLDER"], deck_id)
    meta_path = os.path.join(deck_dir, "meta.txt")
    if not os.path.exists(meta_path):
        return "Deck not found", 404
    with open(meta_path) as f:
        total = int(f.read().strip())

    font_report = {"found": [], "installed": [], "missing": []}
    fonts_path = os.path.join(deck_dir, "fonts.txt")
    if os.path.exists(fonts_path):
        with open(fonts_path) as f:
            for line in f:
                cat, name = line.strip().split(": ", 1)
                font_report[cat].append(name)

    system_fonts = get_system_fonts() if font_report["missing"] else []

    # 加载 slide 内容
    content = []
    content_path = os.path.join(deck_dir, "content.json")
    if os.path.exists(content_path):
        content = load_deck_content(app.config["OUTPUT_FOLDER"], deck_id)

    return render_template(
        "viewer.html",
        deck_id=deck_id,
        total=total,
        font_report=font_report,
        system_fonts=system_fonts,
        slide_content=content,
    )


@app.route("/reconvert/<deck_id>", methods=["POST"])
def reconvert(deck_id):
    deck_dir = os.path.join(app.config["OUTPUT_FOLDER"], deck_id)
    pptx_path = os.path.join(deck_dir, "deck.pptx")
    if not os.path.exists(pptx_path):
        return jsonify({"error": "Deck not found"}), 404

    try:
        replacements = request.json.get("replacements", {})
        weight_map = request.json.get("weights", {})
        if replacements:
            for new_font in replacements.values():
                if not is_font_installed(new_font):
                    try_install_font(new_font)
            subprocess.run(["fc-cache", "-f"], capture_output=True)
            apply_font_replacements(pptx_path, replacements, weight_map)

        convert_deck(deck_dir)

        report = ensure_fonts(pptx_path)
        with open(os.path.join(deck_dir, "fonts.txt"), "w") as f:
            for cat in ("found", "installed", "missing"):
                for name in report[cat]:
                    f.write(f"{cat}: {name}\n")

        return jsonify({"ok": True})
    except Exception as e:
        print(f"[RECONVERT ERROR] {e}", flush=True)
        return jsonify({"error": str(e)}), 500


@app.route("/download/<deck_id>")
def download(deck_id):
    deck_dir = os.path.join(app.config["OUTPUT_FOLDER"], deck_id)
    meta_path = os.path.join(deck_dir, "meta.txt")
    if not os.path.exists(meta_path):
        return "Deck not found", 404
    with open(meta_path) as f:
        total = int(f.read().strip())

    # slide 图片 base64
    slides_data = []
    for i in range(1, total + 1):
        with open(os.path.join(deck_dir, f"slide-{i}.png"), "rb") as f:
            slides_data.append(base64.b64encode(f.read()).decode())

    # slide 内容
    content = []
    content_path = os.path.join(deck_dir, "content.json")
    if os.path.exists(content_path):
        content = load_deck_content(app.config["OUTPUT_FOLDER"], deck_id)

    # 获取当前服务器地址
    server_url = request.host_url.rstrip("/")

    return render_template(
        "download.html",
        deck_id=deck_id,
        total=total,
        slides_data=slides_data,
        slide_content=content,
        server_url=server_url,
    )


@app.route("/slides/<deck_id>/<filename>")
def slide_image(deck_id, filename):
    return send_from_directory(os.path.join(app.config["OUTPUT_FOLDER"], deck_id), filename)


os.makedirs(app.config["OUTPUT_FOLDER"], exist_ok=True)

if __name__ == "__main__":
    socketio.run(app, debug=False, port=8080, host="0.0.0.0", allow_unsafe_werkzeug=True)
