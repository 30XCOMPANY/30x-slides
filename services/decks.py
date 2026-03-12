"""
[INPUT]: 依赖 json、os 与 python-pptx 读取 deck 内容和 viewer 恢复载荷
[OUTPUT]: 对外提供 deck 内容读写、标准化、标题提取、section synopsis 与 system prompt 构建
[POS]: services 的 deck 元数据层，被 app.py 上传流程与 chat.py 对话流程共同消费
[PROTOCOL]: 变更时更新此头部，然后检查 AGENTS.md
"""

import json
import os
import re


IGNORED_TITLES = {"title slide", "thank you", "questions", "end", ""}
GENERIC_THEME_TITLES = {
    "problem", "solution", "roi", "results", "summary", "overview", "agenda",
    "intro", "introduction", "closing", "next steps", "appendix", "why now",
}
ABSTRACT_POSITIONING_TOKENS = {
    "innovation", "innovative", "excellence", "vision", "mission", "leadership",
    "strategy", "values", "value", "culture", "future", "transformation",
    "philosophy", "commitment", "platform", "framework", "principles",
}
STOPWORDS = {
    "the", "and", "for", "with", "from", "that", "this", "into", "about", "your",
    "have", "will", "what", "when", "where", "which", "their", "there", "here",
    "slide", "deck", "overview", "summary", "introduction", "intro", "section",
    "part", "more", "than", "into", "onto", "over", "under", "main", "topic",
}
NAV_INTENT_PATTERNS = (
    "go to", "jump to", "take me to", "show me", "move to", "switch to",
    "next slide", "previous slide", "go back", "next one", "slide ",
    "tell me about", "talk about", "what about", "interested in",
    "want to know", "want to learn", "want to see", "let's look at",
    "let's talk about", "can you explain", "explain the", "more about",
    "dig into", "start with", "cover the", "walk me through",
)
INTERACTIVE_FOLLOW_UPS = (
    "Do you want the short version, the deeper takeaway, or the next slide?",
    "Do you want me to unpack the main point, compare it with the last slide, or move forward?",
    "Should I zoom into the key idea here, connect it to the bigger story, or jump ahead?",
    "Do you want the practical takeaway, the strategy behind it, or the next section?",
    "Want me to simplify this part, go one level deeper, or take you to the next slide?",
)


def normalize_spoken_text(text):
    """Strip visual separators that sound awkward in speech."""
    text = str(text or "")
    text = text.replace("|", ", ")
    text = text.replace("/", ", ")
    text = re.sub(r"\s*,\s*,+", ", ", text)
    text = re.sub(r"\s{2,}", " ", text)
    return text.strip(" ,")


def extract_slide_content(pptx_path):
    """Extract text and speaker notes for each slide."""
    slides_content = []
    try:
        from pptx import Presentation

        prs = Presentation(pptx_path)
        for slide in prs.slides:
            texts = []
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    line = para.text.strip()
                    if line:
                        texts.append(normalize_spoken_text(line))

            notes = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = normalize_spoken_text(slide.notes_slide.notes_text_frame.text.strip())

            slides_content.append({
                "slide": len(slides_content) + 1,
                "text": texts,
                "notes": notes,
            })
    except Exception as exc:
        print(f"[extract] error: {exc}", flush=True)

    return slides_content


def content_path(output_folder, deck_id):
    return os.path.join(output_folder, deck_id, "content.json")


def load_deck_content(output_folder, deck_id):
    path = content_path(output_folder, deck_id)
    if os.path.exists(path):
        with open(path) as handle:
            return json.load(handle)
    return []


def save_deck_content(output_folder, deck_id, content):
    deck_dir = os.path.join(output_folder, deck_id)
    os.makedirs(deck_dir, exist_ok=True)
    with open(content_path(output_folder, deck_id), "w") as handle:
        json.dump(content, handle, ensure_ascii=False)


def normalize_deck_content(raw_content):
    """Accept trusted deck content from the viewer when container-local files are gone."""
    if not isinstance(raw_content, list):
        return []

    normalized = []
    for index, item in enumerate(raw_content, start=1):
        if not isinstance(item, dict):
            continue

        texts = item.get("text", [])
        if not isinstance(texts, list):
            texts = []
        texts = [normalize_spoken_text(str(line).strip()) for line in texts if str(line).strip()]

        notes = item.get("notes", "")
        if not isinstance(notes, str):
            notes = str(notes or "")

        slide_number = item.get("slide", index)
        try:
            slide_number = int(slide_number)
        except Exception:
            slide_number = index

        normalized.append({
            "slide": slide_number,
            "text": texts,
            "notes": normalize_spoken_text(notes.strip()),
        })

    return normalized


def restore_deck_content(output_folder, deck_id, raw_content):
    content = normalize_deck_content(raw_content)
    if content:
        save_deck_content(output_folder, deck_id, content)
        print(f"[TALK] restored {len(content)} slides content from viewer payload", flush=True)
    return content


def get_slide_title(slide_content):
    raw_title = slide_content["text"][0] if slide_content.get("text") else f"Slide {slide_content['slide']}"
    return normalize_spoken_text(raw_title)


def split_phrases(text):
    cleaned = normalize_spoken_text(text)
    if not cleaned:
        return []
    parts = re.split(r"[,:;()\-]\s*|\s+and\s+|\s+vs\.?\s+|\s+to\s+", cleaned)
    return [part.strip() for part in parts if len(part.strip()) >= 3]


def meaningful_tokens(text):
    return [
        token
        for token in re.findall(r"[a-z0-9]+", normalize_spoken_text(text).lower())
        if len(token) >= 3 and token not in STOPWORDS
    ]


def is_weak_theme_phrase(phrase):
    normalized = normalize_spoken_text(phrase)
    lowered = normalized.lower()
    if lowered in IGNORED_TITLES or lowered in GENERIC_THEME_TITLES:
        return True
    if re.search(r"\b20\d{2}\b", lowered):
        return True
    if len(meaningful_tokens(lowered)) <= 1:
        return True
    return False


def is_abstract_positioning_phrase(phrase):
    tokens = meaningful_tokens(phrase)
    if not tokens:
        return True
    abstract_hits = sum(1 for token in tokens if token in ABSTRACT_POSITIONING_TOKENS)
    return abstract_hits >= max(1, len(tokens) - 1)


def extract_theme_phrases(all_content, limit=4):
    scores = {}
    ordered = []

    for idx, slide_content in enumerate(all_content):
        title = get_slide_title(slide_content)
        title_phrases = split_phrases(title) or ([title] if title else [])
        body_phrases = []
        for line in slide_content.get("text", [])[1:4]:
            body_phrases.extend(split_phrases(line)[:1])

        for phrase in title_phrases[:2] + body_phrases[:2]:
            normalized = phrase.lower()
            if len(normalized) < 3 or is_weak_theme_phrase(phrase):
                continue
            if normalized not in scores:
                ordered.append(normalized)
                scores[normalized] = {"phrase": phrase, "score": 0}
            weight = 4 if idx == 0 else 2
            if phrase in body_phrases:
                weight = 3
            scores[normalized]["score"] += weight

    ranked = sorted(ordered, key=lambda key: (-scores[key]["score"], ordered.index(key)))
    return [scores[key]["phrase"] for key in ranked[:limit]]


def join_phrases(phrases):
    if not phrases:
        return ""
    if len(phrases) == 1:
        return phrases[0]
    if len(phrases) == 2:
        return f"{phrases[0]} and {phrases[1]}"
    return ", ".join(phrases[:-1]) + f", and {phrases[-1]}"


def interactive_follow_up(seed_text="", slide_number=0):
    seed = f"{slide_number}:{normalize_spoken_text(seed_text).lower()}"
    index = sum(ord(ch) for ch in seed) % len(INTERACTIVE_FOLLOW_UPS)
    return INTERACTIVE_FOLLOW_UPS[index]


def summarize_slide_core(slide_content):
    title = get_slide_title(slide_content)
    body_lines = [line for line in slide_content.get("text", []) if line.strip()]
    detail_lines = body_lines[1:5] if len(body_lines) > 1 else body_lines[:4]

    concept_phrases = []
    seen = set()
    for line in detail_lines:
        for phrase in split_phrases(line)[:2]:
            normalized = phrase.lower()
            if normalized == title.lower() or normalized in seen:
                continue
            seen.add(normalized)
            concept_phrases.append(phrase)
            if len(concept_phrases) >= 3:
                break
        if len(concept_phrases) >= 3:
            break

    return {
        "title": title,
        "concepts": concept_phrases,
        "body_lines": detail_lines,
    }


def choose_synopsis_section_count(total_slides):
    if total_slides >= 18:
        return 5
    if total_slides >= 6:
        return 4
    return max(4, min(5, total_slides))


def partition_slide_ranges(total_slides, section_count):
    if total_slides <= 0 or section_count <= 0:
        return []

    base = total_slides // section_count
    remainder = total_slides % section_count
    ranges = []
    start = 0
    for idx in range(section_count):
        size = base + (1 if idx < remainder else 0)
        if size <= 0:
            continue
        end = start + size
        ranges.append((start, end))
        start = end
    return ranges


def best_section_phrase(section_slides):
    phrase_scores = {}
    phrase_order = []

    for offset, slide_content in enumerate(section_slides):
        summary = summarize_slide_core(slide_content)
        title = summary["title"]
        title_phrases = split_phrases(title) or ([title] if title else [])
        body_phrases = []
        for line in summary["body_lines"][:3]:
            body_phrases.extend(split_phrases(line)[:2])

        for phrase in title_phrases[:3] + body_phrases[:3]:
            normalized = phrase.lower()
            if is_weak_theme_phrase(phrase):
                continue
            if normalized not in phrase_scores:
                phrase_scores[normalized] = {"phrase": phrase, "score": 0}
                phrase_order.append(normalized)

            score = 5 if phrase in title_phrases else 3
            if offset == 0:
                score += 1
            if is_abstract_positioning_phrase(phrase):
                score -= 3
            if len(meaningful_tokens(phrase)) >= 3:
                score += 1
            phrase_scores[normalized]["score"] += score

    if not phrase_scores:
        for slide_content in section_slides:
            title = get_slide_title(slide_content)
            if not is_weak_theme_phrase(title):
                return title
        return "the next part of the story"

    ranked = sorted(
        phrase_order,
        key=lambda key: (-phrase_scores[key]["score"], phrase_order.index(key))
    )
    return phrase_scores[ranked[0]]["phrase"]


def best_section_detail(section_slides, section_label):
    detail_scores = {}
    detail_order = []

    for slide_content in section_slides:
        summary = summarize_slide_core(slide_content)
        for phrase in summary["concepts"] + summary["body_lines"][:3]:
            normalized = phrase.lower()
            if (
                not normalized
                or normalized == section_label.lower()
                or is_weak_theme_phrase(phrase)
                or is_abstract_positioning_phrase(phrase)
            ):
                continue

            if normalized not in detail_scores:
                detail_scores[normalized] = {"phrase": phrase, "score": 0}
                detail_order.append(normalized)

            score = 4 if phrase in summary["concepts"] else 2
            if len(meaningful_tokens(phrase)) >= 3:
                score += 1
            detail_scores[normalized]["score"] += score

    if not detail_scores:
        return ""

    ranked = sorted(
        detail_order,
        key=lambda key: (-detail_scores[key]["score"], detail_order.index(key))
    )
    return detail_scores[ranked[0]]["phrase"]


def build_section_spoken_concept(section_label, section_detail):
    label = normalize_spoken_text(section_label)
    detail = normalize_spoken_text(section_detail)

    if detail and detail.lower() != label.lower():
        return f"{label}, especially {detail}"
    return label


def build_deck_synopsis(all_content):
    total_slides = len(all_content)
    if not total_slides:
        return []

    section_count = min(choose_synopsis_section_count(total_slides), total_slides)
    synopsis = []
    for start, end in partition_slide_ranges(total_slides, section_count):
        section_slides = all_content[start:end]
        if not section_slides:
            continue
        label = best_section_phrase(section_slides)
        synopsis.append({
            "label": label,
            "spoken_concept": build_section_spoken_concept(
                label,
                best_section_detail(section_slides, label),
            ),
            "start_slide": section_slides[0]["slide"],
            "end_slide": section_slides[-1]["slide"],
        })
    return synopsis


def build_navigation_index(all_content):
    index = []
    for slide_content in all_content:
        title = get_slide_title(slide_content)
        summary = summarize_slide_core(slide_content)
        phrases = [title] + summary["concepts"]
        body_lines = summary["body_lines"]
        if body_lines:
            phrases.extend(body_lines[:2])

        keywords = set()
        for phrase in phrases:
            keywords.update(meaningful_tokens(phrase))

        index.append({
            "slide": slide_content["slide"],
            "title": title,
            "keywords": keywords,
            "summary": summary,
        })
    return index


def infer_navigation_target(user_text, all_content, current_slide_idx):
    normalized = normalize_spoken_text(user_text).lower()
    if not normalized:
        return None

    if "next slide" in normalized or normalized in {"next", "next one", "move on"}:
        return "next"
    if "previous slide" in normalized or "go back" in normalized or normalized in {"back", "previous", "prev"}:
        return "prev"

    requested_number = re.search(r"\bslide\s+(\d{1,2})\b", normalized)
    if requested_number:
        return requested_number.group(1)

    has_nav_intent = any(pattern in normalized for pattern in NAV_INTENT_PATTERNS)
    if not has_nav_intent:
        return None

    nav_index = build_navigation_index(all_content)
    query_tokens = meaningful_tokens(normalized)
    if not query_tokens:
        return None

    best_slide = None
    best_score = 0
    current_slide_number = current_slide_idx + 1
    for candidate in nav_index:
        phrase_match = False
        candidate_phrases = [candidate["title"]] + candidate["summary"]["concepts"] + candidate["summary"]["body_lines"][:2]
        for phrase in candidate_phrases:
            lowered = normalize_spoken_text(phrase).lower()
            if lowered and lowered in normalized:
                phrase_match = True
                break

        overlap = len(candidate["keywords"].intersection(query_tokens))
        if overlap <= 0 and not phrase_match:
            continue
        score = overlap * 3
        if phrase_match:
            score += 5
        if candidate["slide"] == current_slide_number:
            score -= 1
        title_tokens = set(meaningful_tokens(candidate["title"]))
        if title_tokens.intersection(query_tokens):
            score += 2
        if score > best_score:
            best_score = score
            best_slide = candidate["slide"]

    return str(best_slide) if best_slide and best_score >= 3 and best_slide != current_slide_number else None


def build_greeting(all_content):
    """Template fallback greeting — used only when LLM is unavailable."""
    synopsis = build_deck_synopsis(all_content)
    if not synopsis:
        return (
            "Hey, glad you're here. I can walk you through this deck. "
            "What do you want to know?"
        )

    section_labels = [section["label"] for section in synopsis[:5]]
    joined = join_phrases(section_labels)

    return (
        f"Hey, glad you're here. This deck covers {joined}. "
        "I can break any of that down for you. What do you want to start with?"
    )


def build_greeting_prompt(all_content):
    """Build a system prompt specifically for generating the greeting with deck summary."""
    slide_overview = ""
    for sc in all_content:
        title = get_slide_title(sc)
        body = " | ".join(sc.get("text", [])[:3])
        slide_overview += f"  Slide {sc['slide']}: {title} — {body}\n"

    return f"""You are a friendly guide helping someone explore a slide deck. Your job:

1. Say hi naturally (one short sentence, like "Hey, glad you're here.").
2. Give a 2-3 sentence SUMMARY of what this deck is about — the story, the argument, the key message. Synthesize, don't list slide titles. Talk about it like you're telling a friend what this presentation covers.
3. End with one short question asking what they want to explore first.

RULES:
- Total length: 4-5 sentences max.
- Spoken voice only — no markdown, no bullets.
- You are NOT the speaker or presenter. You are a guide helping the viewer understand the deck.
- Never say "I'm the speaker" or "as the presenter" or anything like that.
- Never greet twice. This greeting happens exactly once.

DECK CONTENT:
{slide_overview}"""


def build_local_slide_reply(all_content, current_slide_idx, user_text=""):
    """Fallback explanation when the hosted LLM is unavailable."""
    if not all_content:
        return "I don't have the slide content loaded yet. Want me to try again?"

    if current_slide_idx < 0 or current_slide_idx >= len(all_content):
        current_slide_idx = 0

    slide_content = all_content[current_slide_idx]
    summary = summarize_slide_core(slide_content)
    title = summary["title"]
    concepts = summary["concepts"]
    body_lines = summary["body_lines"]

    if not body_lines:
        return (
            f"So {title} is really the setup here. "
            "Want me to go deeper or move on?"
        )

    concept_text = join_phrases(concepts[:3])
    parts = []

    if concept_text:
        parts.append(f"So the big idea here is {concept_text}.")
    else:
        parts.append(f"So {body_lines[0]}.")

    if len(body_lines) > 1:
        parts.append(f"And then {body_lines[1]}.")

    parts.append("Want me to unpack anything here, or move on?")
    return " ".join(parts[:3])


def build_system_prompt(all_content, current_slide_idx, allow_control_tags=True):
    title_index = ""
    for slide_content in all_content:
        title = get_slide_title(slide_content)
        title_index += f"  Slide {slide_content['slide']}: {title}\n"

    current_detail = ""
    if 0 <= current_slide_idx < len(all_content):
        slide_content = all_content[current_slide_idx]
        current_detail = f"Slide {slide_content['slide']} — {get_slide_title(slide_content)}:\n"
        current_detail += "\n".join(slide_content["text"])
        if slide_content.get("notes"):
            current_detail += f"\nSpeaker notes: {slide_content['notes']}"

    nav_block = ""
    if allow_control_tags:
        nav_examples = ""
        example_slides = []
        for slide_content in all_content:
            title = get_slide_title(slide_content)
            if title and title.lower() not in IGNORED_TITLES:
                example_slides.append((slide_content["slide"], title))

        for slide_number, slide_title in example_slides[:3]:
            keyword = slide_title.split()[0].lower() if slide_title.split() else slide_title.lower()
            nav_examples += (
                f'- User says "{keyword}" → you write [GO:{slide_number}] '
                f'because Slide {slide_number} is "{slide_title}"\n'
            )
        nav_examples += '- User says "next" → you write [GO:next]\n'
        nav_examples += '- User says "go back" → you write [GO:prev]\n'

        nav_block = f"""
===== NAVIGATION RULES (MANDATORY) =====
When the user mentions ANY topic, keyword, or phrase that relates to a slide title, you MUST include [GO:N].

EXAMPLES from this deck:
{nav_examples}
HOW IT WORKS: You write [GO:N] anywhere in your reply. The system removes it before showing to user and auto-jumps the slide.

RULES:
1. Match loosely. Any recognizable keyword should navigate.
2. Never ask whether to navigate. Just include [GO:N].
3. "next" → [GO:next], "back"/"previous" → [GO:prev].
4. If user asks about a topic and you omit [GO:N], your response is wrong.
5. Always include [GO:N] before your spoken text, like: [GO:3] So this one covers...
======================================="""
    else:
        nav_block = """
NAVIGATION:
Do not output control tags or bracketed commands.
Answer naturally only. Navigation is handled outside the model."""

    return f"""You are a friendly guide helping someone explore a slide deck.

RULES:
- Spoken voice only. No markdown, bullets, or lists.
- Casual and warm. Like a smart friend explaining a topic.
- NEVER say "this slide mentions", "this slide shows", "this slide covers", "the slide talks about", or any variation. Just dive straight into the content. Instead of "This slide covers customer challenges", say "So the big customer challenge here is..."
- Synthesize the ideas, don't narrate the slide. Explain WHY it matters.
- Keep replies to 3-5 short sentences. End with a natural question or offer.
- Never repeat what you just said in the same turn.
- You are NOT the speaker or presenter. Never say "I" when referring to the deck's author.
- If the user says hi/hello and you already greeted them, don't greet again — just respond naturally.

Viewer is on Slide {current_slide_idx + 1} of {len(all_content)}.

SLIDE TITLES:
{title_index}
CURRENT SLIDE:
{current_detail}
{nav_block}"""
